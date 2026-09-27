from flask import Blueprint, abort, render_template, request, send_file
from platform_core.services.workspace_service import workspace_context
from .services.customs_service import lookup_tariffs,load_customs,calculate_exchange,rates_dataset

bp=Blueprint("tab2_customs",__name__,template_folder="templates",static_folder="static",static_url_path="/customs-assets")

@bp.route("/beauty/<country>/customs",methods=["GET","POST"])
def index(country):
    try:
        ctx=workspace_context(country)
    except ValueError:
        abort(404)
    product=ctx["selected_product"] or {}
    hs=request.values.get("hs_code",product.get("hs_code",""))
    origin=request.values.get("origin","KR").upper()
    result=None
    exchange=None
    error=None
    try:
        if request.method=="POST" and request.form.get("action")=="exchange":
            exchange=calculate_exchange(request.form.get("amount",""),request.form.get("from_currency","KRW"),request.form.get("to_currency",ctx["destination"]["currency"]))
        if hs:
            result=lookup_tariffs(country,origin,hs)
    except ValueError as exc:
        error=str(exc)
    return render_template("tab2_customs/index.html",**ctx,tab="customs",result=result,hs_code=hs,origin=origin,error=error,exchange=exchange,customs=load_customs(country),rates=rates_dataset(),available_currencies=sorted({"KRW",ctx["destination"]["currency"]}|{r["currency"] for r in rates_dataset()["data"].get("rates",[])}))

@bp.get("/beauty/<country>/customs/export")
def export_excel(country):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    try:
        ctx = workspace_context(country)
    except ValueError:
        abort(404)
    if not ctx["project"]:
        abort(400, description="내보낼 프로젝트를 선택해 주세요.")
    product = ctx["selected_product"] or {}
    hs = request.args.get("hs_code", product.get("hs_code", ""))
    origin = request.args.get("origin", "KR").upper()
    try:
        tariffs = lookup_tariffs(country, origin, hs) if hs else {
            "items": [], "message": "HS 코드를 입력해 주세요."
        }
    except ValueError as exc:
        tariffs = {"items": [], "message": str(exc)}
    customs = load_customs(country)
    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(title, headers, rows):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                # 사용자 입력과 출처 문자열을 엑셀 수식으로 실행하지 않습니다.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 26

    add_sheet("프로젝트 정보", ["항목", "내용"], [
        ["프로젝트", ctx["project"]["name"]],
        ["목적국", ctx["destination"]["name"]],
        ["선택 제품", product.get("name", "선택 제품 없음")],
        ["원산지", origin],
        ["조회 HS 코드", hs],
        ["관세 안내", "FTA 세율은 원산지 요건과 증빙을 충족할 때 적용 가능합니다. 6자리 코드 조회 시 목적국 세부 품목을 확인하세요."],
    ])
    tariff_rows = [
        [
            item.get("product_name") or "품목명 미등록",
            str(item.get("hs_code", "")),
            item.get("rate_type") or "미확인",
            item.get("rate") if item.get("rate") is not None else "미확인",
            item.get("rate_unit") or "",
            item.get("agreement") or "협정 미등록",
            item.get("conditions") or "원산지 요건·증빙 확인 필요",
            item.get("as_of") or "기준일 미등록",
            item.get("source") or "",
            item.get("source_url") or "",
        ] for item in tariffs["items"]
    ]
    add_sheet("관세율", [
        "품목", "HS 코드", "세율 유형", "관세율", "단위", "협정",
        "적용 조건", "기준일", "출처", "출처 URL",
    ], tariff_rows or [[tariffs.get("message") or "조회 조건에 맞는 등록 데이터가 없습니다."]])
    customs_rows = []
    for number, step in enumerate(customs["steps"], 1):
        documents = step.get("required_doc") or "미등록"
        if isinstance(documents, list):
            documents = " · ".join(str(document) for document in documents)
        customs_rows.append([
            number, step.get("task_name") or step.get("title") or "",
            step.get("description") or "", documents,
            step.get("owner") or "미등록", step.get("guide_url") or "",
        ])
    add_sheet("통관 안내", [
        "단계", "업무", "설명", "필요 서류", "담당 주체", "공식 안내 URL",
    ], customs_rows or [[customs.get("message") or "국가별 통관 안내 준비 중"]])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f'beauty-customs-{ctx["project"]["id"][:8]}.xlsx',
    )
